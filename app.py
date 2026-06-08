import streamlit as st
import pandas as pd
import re
import io
@st.cache_data
def carregar_glossario():

    df = pd.read_csv(
        "Glossário_2021.csv",
        sep=";",
        encoding="utf-8-sig",
        header=2,
        engine="python",
        on_bad_lines="skip"
    )

    df.columns = [
        str(c).strip()
        for c in df.columns
    ]

    return df
from difflib import SequenceMatcher
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =====================================================
# CONFIGURAÇÃO
# =====================================================

st.set_page_config(
    page_title="Minerador de Termos ANEEL",
    layout="wide"
)

st.title("📚 Minerador de Termos para Glossário ANEEL")

# =====================================================
# FUNÇÕES
# =====================================================

TERMO_COL = "Termo"
DEFINICAO_COL = "Definição"


def limpar_excel(valor):
    if pd.isna(valor):
        return valor

    valor = str(valor)
    return ILLEGAL_CHARACTERS_RE.sub("", valor)


def similaridade(a, b):
    return SequenceMatcher(
        None,
        str(a).lower(),
        str(b).lower()
    ).ratio()


# =====================================================
# LEITURA GLOSSÁRIO
# =====================================================

def ler_glossario(arquivo):

    nome = arquivo.name.lower()

    try:

        if nome.endswith(".csv"):

            df = pd.read_csv(
                arquivo,
                sep=";",
                encoding="utf-8-sig",
                header=2,
                engine="python",
                on_bad_lines="skip"
            )

        elif nome.endswith(".xlsx"):

            df = pd.read_excel(arquivo)

        elif nome.endswith(".xls"):

            df = pd.read_excel(arquivo)

        elif nome.endswith(".ods"):

            df = pd.read_excel(
                arquivo,
                engine="odf"
            )

        else:
            st.error("Formato de glossário não suportado.")
            return None

        df.columns = [
            str(c).strip()
            for c in df.columns
        ]

        if TERMO_COL not in df.columns:
            st.error(
                f"Coluna '{TERMO_COL}' não encontrada."
            )
            st.write(df.columns.tolist())
            return None

        if DEFINICAO_COL not in df.columns:
            st.error(
                f"Coluna '{DEFINICAO_COL}' não encontrada."
            )
            st.write(df.columns.tolist())
            return None

        return df

    except Exception as e:
        st.error(f"Erro ao ler glossário: {e}")
        return None


# =====================================================
# LEITURA DE DOCUMENTOS
# =====================================================

def ler_pdf(arquivo):

    try:

        from pypdf import PdfReader

        reader = PdfReader(arquivo)

        texto = ""

        for pagina in reader.pages:

            conteudo = pagina.extract_text()

            if conteudo:
                texto += conteudo + "\n"

        return texto

    except Exception as e:
        st.error(f"Erro PDF: {e}")
        return ""


def ler_txt(arquivo):

    try:
        return arquivo.read().decode("utf-8", errors="ignore")
    except:
        return ""


def ler_html(arquivo):

    try:

        from bs4 import BeautifulSoup

        html = arquivo.read().decode(
            "utf-8",
            errors="ignore"
        )

        soup = BeautifulSoup(
            html,
            "html.parser"
        )

        return soup.get_text(" ")

    except Exception as e:
        st.error(f"Erro HTML: {e}")
        return ""


def obter_texto(arquivo):

    nome = arquivo.name.lower()

    if nome.endswith(".pdf"):
        return ler_pdf(arquivo)

    if nome.endswith(".txt"):
        return ler_txt(arquivo)

    if nome.endswith(".html"):
        return ler_html(arquivo)

    if nome.endswith(".htm"):
        return ler_html(arquivo)

    return ""


# =====================================================
# EXTRAÇÃO DE TERMOS
# =====================================================

def extrair_candidatos(texto):

    candidatos = []

    padrao = re.compile(
        r'([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][A-Za-zÁÀÂÃÉÊÍÓÔÕÚÇáàâãéêíóôõúç0-9\s\-]{3,80})'
    )

    encontrados = padrao.findall(texto)

    vistos = set()

    for termo in encontrados:

        termo = " ".join(
            termo.split()
        ).strip()

        if len(termo) < 4:
            continue

        if termo.lower() in vistos:
            continue

        vistos.add(
            termo.lower()
        )

        candidatos.append(termo)

    return candidatos


# =====================================================
# COMPARAÇÃO
# =====================================================

def avaliar_termo(
    termo,
    indice_glossario,
    termos_existentes
):

    termo_norm = (
        str(termo)
        .lower()
        .strip()
    )

    # ==========================
    # termo já existe
    # ==========================

    if termo_norm in termos_existentes:

        return (
            100,
            "verificar atualização",
            termo
        )

    melhor_sim = 0
    melhor_termo = ""

    letra = termo_norm[0]

    candidatos = indice_glossario.get(
        letra,
        []
    )

    tamanho = len(termo)

    for termo_glossario in candidatos:

        # filtro por tamanho

        if abs(
            len(str(termo_glossario))
            - tamanho
        ) > 10:
            continue

        sim = similaridade(
            termo,
            termo_glossario
        )

        if sim > melhor_sim:

            melhor_sim = sim
            melhor_termo = termo_glossario

    score = round(
        melhor_sim * 100,
        2
    )

    if score >= 95:

        classificacao = (
            "verificar atualização"
        )

    elif score >= 75:

        classificacao = "Revisar"

    elif score >= 50:

        classificacao = (
            "Fraco candidato"
        )

    else:

        classificacao = (
            "Forte candidato"
        )

    return (
        score,
        classificacao,
        melhor_termo
    )


# =====================================================
# INTERFACE
# =====================================================

st.subheader("1. Upload dos documentos")

documentos = st.file_uploader(
    "Arquivos DOU",
    type=[
        "pdf",
        "txt",
        "html",
        "htm"
    ],
    accept_multiple_files=True
)

# =====================================================
# PROCESSAMENTO
# =====================================================

if documentos:

    with st.spinner("Carregando glossário..."):

        glossario_df = carregar_glossario()

    if glossario_df is not None:

        st.success(
            f"Glossário carregado: {len(glossario_df)} registros"
        )

        # ==========================
        # Índice rápido
        # ==========================

        glossario_df["termo_norm"] = (
            glossario_df[TERMO_COL]
            .astype(str)
            .str.lower()
            .str.strip()
        )

        termos_existentes = set(
            glossario_df["termo_norm"]
        )

        indice_glossario = {}

        for termo_glossario in glossario_df[TERMO_COL].dropna():

            termo_glossario = str(
                termo_glossario
            ).strip()

            if not termo_glossario:
                continue

            letra = termo_glossario[0].lower()

            indice_glossario.setdefault(
                letra,
                []
            ).append(
                termo_glossario
            )

        resultados = []

        for doc in documentos:

            texto = obter_texto(doc)

            candidatos = extrair_candidatos(
                texto
            )

            progress = st.progress(0)

            total = len(candidatos)

            for i, termo in enumerate(candidatos):

                progress.progress((i + 1) / total)

                (
                    score,
                    classificacao,
                    termo_existente
                ) = avaliar_termo(
                    termo,
                    indice_glossario,
                    termos_existentes
                )

                resultados.append(
                    {
                        "Arquivo": doc.name,
                        "Termo candidato": termo,
                        "Similaridade (%)": score,
                        "Termo mais próximo": termo_existente,
                        "Classificação": classificacao
                    }
                )

        st.write(
            f"Resultados encontrados: {len(resultados)}"
        )

        resultado_df = pd.DataFrame(
            resultados
        )

        if not resultado_df.empty:

            resultado_df = resultado_df.sort_values(
                by="Similaridade (%)",
                ascending=False
            )

            for coluna in resultado_df.columns:
                resultado_df[coluna] = resultado_df[coluna].apply(
                    limpar_excel
                )

            st.subheader(
                "Resultados"
            )

            st.dataframe(
                resultado_df,
                use_container_width=True
            )

            buffer = io.BytesIO()

            with pd.ExcelWriter(
                buffer,
                engine="openpyxl"
            ) as writer:

                resultado_df.to_excel(
                    writer,
                    index=False,
                    sheet_name="Resultados"
                )

                ws = writer.sheets["Resultados"]
        
# Cabeçalho
for cell in ws[1]:
    cell.font = Font(
        bold=True,
        color="FFFFFF"
    )

    cell.fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

# Ajustar largura das colunas
for coluna in ws.columns:

    tamanho = max(
        len(str(cell.value))
        if cell.value is not None
        else 0
        for cell in coluna
    )

    letra = get_column_letter(
        coluna[0].column
    )

    ws.column_dimensions[
        letra
    ].width = min(
        tamanho + 3,
        60
    )
    
    st.download_button(
                "📥 Baixar Excel",
                data=buffer.getvalue(),
                file_name="resultado_glossario.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
else:
            st.warning(
                "Nenhum candidato encontrado."
            )
